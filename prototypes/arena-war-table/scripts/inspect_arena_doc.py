from pathlib import Path
from docx import Document
from docx.oxml.ns import qn


DOC_PATH = Path(r"C:\Projects\Wildcrown\Wildcrown_Arena_and_9v3_Coop_Design.docx")


def main() -> None:
    document = Document(DOC_PATH)
    print(f"paragraphs={len(document.paragraphs)} inline_shapes={len(document.inline_shapes)}")
    for index, paragraph in enumerate(document.paragraphs):
        drawings = paragraph._p.xpath(".//w:drawing")
        text = paragraph.text.strip().replace("\t", " ")
        if drawings or text:
            print(f"{index:03d} drawings={len(drawings)} style={paragraph.style.name!r} text={text[:180]!r}")
            for drawing in drawings:
                blips = drawing.xpath(".//a:blip")
                for blip in blips:
                    rel_id = blip.get(qn("r:embed"))
                    rel = document.part.rels.get(rel_id)
                    print(f"    image_rel={rel_id} target={getattr(rel, 'target_ref', None)}")


if __name__ == "__main__":
    main()

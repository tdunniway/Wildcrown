from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


SOURCE_DOC = Path(r"C:\Projects\Wildcrown\Wildcrown_Arena_and_9v3_Coop_Design.docx")
IMAGE_PATH = Path(r"C:\Projects\Wildcrown\prototypes\arena-war-table\public\assets\arena-war-table-reference.png")
OUTPUT_DOC = Path(r"C:\Projects\Wildcrown\prototypes\arena-war-table\deliverables\Wildcrown_Arena_and_9v3_Coop_Design.docx")
ANCHOR_TEXT = "Design handoff. Version 1.0"
CAPTION_TEXT = "Arena War Table prototype. Selected UI direction 1."


def set_picture_alt_text(inline_shape, title: str, description: str) -> None:
    doc_properties = inline_shape._inline.docPr
    doc_properties.set("title", title)
    doc_properties.set("descr", description)


def main() -> None:
    document = Document(SOURCE_DOC)
    anchor = next(
        (paragraph for paragraph in document.paragraphs if paragraph.text.startswith(ANCHOR_TEXT)),
        None,
    )
    if anchor is None:
        raise RuntimeError(f"Could not find insertion anchor: {ANCHOR_TEXT}")

    image_paragraph = anchor.insert_paragraph_before()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.space_before = Pt(8)
    image_paragraph.paragraph_format.space_after = Pt(4)
    image_run = image_paragraph.add_run()
    inline_shape = image_run.add_picture(str(IMAGE_PATH), width=Inches(7.0))
    set_picture_alt_text(
        inline_shape,
        "Wildcrown Arena War Table interface",
        "Landscape planning screen with four champion cards, three tactical lanes, opening orders, and a commit plan control.",
    )

    caption = anchor.insert_paragraph_before(CAPTION_TEXT, style="Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(10)

    OUTPUT_DOC.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_DOC)
    print(OUTPUT_DOC)


if __name__ == "__main__":
    main()
